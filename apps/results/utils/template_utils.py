from openpyxl import Workbook

from apps.accounts.models import User

from .excel_constants import (
    COLUMN_MARKS_OBTAINED,
    COLUMN_MAX_MARKS,
    DEFAULT_MAX_MARKS,
    HEADER_ROW,
    IMPORT_HEADERS,
)
from .excel_styles import (
    EDITABLE_FILL,
    LOCKED,
    NORMAL_FONT,
    TITLE_FONT,
    THIN_BORDER,
    UNLOCKED,
    auto_fit_columns,
    freeze_header,
    protect_sheet,
    style_data_row,
    style_header_row,
    CENTER,
)


def generate_result_import_template(
    *,
    school_class,
    exam,
):
    """
    Generate a multi-sheet Excel workbook.

    One worksheet is created for each subject
    assigned to the class.

    Teachers only edit the
    'Marks Obtained' column.
    """

    workbook = Workbook()

    workbook.remove(workbook.active)

    students = (
        User.objects.filter(
            role=User.Role.STUDENT,
            school=school_class.school,
            school_class=school_class,
        )
        .order_by(
            "username",
            "first_name",
            "last_name",
        )
    )

    subjects = school_class.subjects.order_by("name")

    for subject in subjects:

        sheet = workbook.create_sheet(
            title=subject.name[:31],
        )

        # ----------------------------------------
        # Title
        # ----------------------------------------

        sheet.merge_cells("A1:D1")

        cell = sheet["A1"]

        cell.value = school_class.school.name
        cell.font = TITLE_FONT
        cell.alignment = CENTER

        sheet["A2"] = "Exam"
        sheet["B2"] = exam.display_name

        sheet["A3"] = "Class"
        sheet["B3"] = (
            f"{school_class.name} "
            f"{school_class.section}"
        ).strip()

        sheet["C2"] = "Subject"
        sheet["D2"] = subject.name

        sheet["C3"] = "Subject Code"
        sheet["D3"] = subject.code

        sheet["A4"] = "Academic Year"
        sheet["B4"] = exam.academic_year

        # ----------------------------------------
        # Header
        # ----------------------------------------

        for column, header in enumerate(
            IMPORT_HEADERS,
            start=1,
        ):
            sheet.cell(
                row=HEADER_ROW,
                column=column,
                value=header,
            )

        style_header_row(
            sheet,
            HEADER_ROW,
        )

        # ----------------------------------------
        # Students
        # ----------------------------------------

        row = HEADER_ROW + 1

        for student in students:

            sheet.cell(
                row=row,
                column=1,
                value=student.username,
            )

            sheet.cell(
                row=row,
                column=2,
                value=student.get_full_name(),
            )

            marks_cell = sheet.cell(
                row=row,
                column=COLUMN_MARKS_OBTAINED,
            )

            marks_cell.fill = EDITABLE_FILL
            marks_cell.protection = UNLOCKED

            sheet.cell(
                row=row,
                column=COLUMN_MAX_MARKS,
                value=DEFAULT_MAX_MARKS,
            )

            for cell in sheet[row]:
                cell.font = NORMAL_FONT
                cell.border = THIN_BORDER

            row += 1

        # ----------------------------------------
        # Protect Readonly Cells
        # ----------------------------------------

        for rows in sheet.iter_rows(
            min_row=HEADER_ROW + 1,
            max_row=row - 1,
        ):

            style_data_row(rows)

            for cell in rows:

                if cell.column != COLUMN_MARKS_OBTAINED:
                    cell.protection = LOCKED

        freeze_header(
            sheet,
            HEADER_ROW,
        )

        auto_fit_columns(
            sheet,
        )

        protect_sheet(
            sheet,
        )

    return workbook