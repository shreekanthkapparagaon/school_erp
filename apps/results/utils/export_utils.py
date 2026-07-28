import re
from io import BytesIO

from django.http import HttpResponse

from openpyxl import Workbook

from .excel_constants import (
    CLASS_EXPORT_HEADERS,
    HEADER_ROW,
    SCHOOL_EXPORT_HEADERS,
)
from .excel_styles import (
    auto_fit_columns,
    freeze_header,
    style_data_row,
    style_header_row,
)


def safe_sheet_title(
    value,
    used_titles,
):
    """
    Generate a valid, unique Excel worksheet title.
    """

    title = re.sub(
        r"[\[\]:*?/\\]",
        "-",
        value,
    ).strip()[:31] or "Sheet"

    original = title
    counter = 2

    while title in used_titles:
        suffix = f"-{counter}"
        title = f"{original[:31 - len(suffix)]}{suffix}"
        counter += 1

    used_titles.add(title)
    return title


def append_results_sheet(
    *,
    workbook,
    title,
    headers,
    results,
    used_titles,
):
    """
    Create a worksheet containing the supplied results.
    """

    sheet = workbook.create_sheet(
        title=safe_sheet_title(
            title,
            used_titles,
        )
    )

    # Write header row
    for column, header in enumerate(
        headers,
        start=1,
    ):
        sheet.cell(
            row=HEADER_ROW,
            column=column,
            value=header,
        )

    style_header_row(
        sheet,
        HEADER_ROW,
    )

    row = HEADER_ROW + 1

    is_class_export = headers == CLASS_EXPORT_HEADERS

    for result in results:

        if is_class_export:
            row_data = [
                str(result.student),
                result.student.email,
                float(result.marks_obtained),
                float(result.max_marks),
                float(result.percentage),
                result.grade,
                "Yes" if result.is_published else "No",
            ]
        else:
            row_data = [
                str(result.student),
                result.student.email,
                result.subject.name,
                result.subject.code,
                float(result.marks_obtained),
                float(result.max_marks),
                float(result.percentage),
                result.grade,
                result.exam.name,
                result.exam.academic_year or "",
                "Yes" if result.is_published else "No",
            ]

        sheet.append(row_data)

        style_data_row(
            sheet[row]
        )

        row += 1

    freeze_header(
        sheet,
        HEADER_ROW,
    )

    auto_fit_columns(
        sheet,
    )


def create_class_results_workbook(
    *,
    school_class,
    results,
):
    """
    Create an Excel workbook for a class.

    Each subject is exported into its own worksheet.
    """

    workbook = Workbook()
    workbook.remove(workbook.active)

    used_titles = set()
    subject_results = {}

    for result in results:
        subject_results.setdefault(
            result.subject,
            [],
        ).append(result)

    if not subject_results:
        sheet = workbook.create_sheet("No Results")
        sheet["A1"] = (
            f"No results have been entered for "
            f"{school_class}."
        )
        return workbook

    for subject in sorted(
        subject_results.keys(),
        key=lambda subject: subject.name.lower(),
    ):
        append_results_sheet(
            workbook=workbook,
            title=subject.name,
            headers=CLASS_EXPORT_HEADERS,
            results=subject_results[subject],
            used_titles=used_titles,
        )

    return workbook


def create_school_results_workbook(
    *,
    exam,
    classes,
):
    """
    Create an Excel workbook containing every class for an
    examination.

    Each worksheet represents one class and one subject.
    """

    workbook = Workbook()
    workbook.remove(workbook.active)

    used_titles = set()

    for school_class in classes:

        results = (
            school_class.results.filter(
                exam=exam,
            )
            .select_related(
                "student",
                "subject",
                "exam",
            )
            .order_by(
                "subject__name",
                "student__first_name",
                "student__last_name",
            )
        )

        if not results.exists():
            continue

        subject_results = {}

        for result in results:
            subject_results.setdefault(
                result.subject,
                [],
            ).append(result)

        for subject in sorted(
            subject_results.keys(),
            key=lambda subject: subject.name.lower(),
        ):
            append_results_sheet(
                workbook=workbook,
                title=f"{school_class} - {subject.name}",
                headers=SCHOOL_EXPORT_HEADERS,
                results=subject_results[subject],
                used_titles=used_titles,
            )

    if not workbook.sheetnames:
        sheet = workbook.create_sheet("No Results")
        sheet["A1"] = (
            f"No results have been entered for "
            f"'{exam.name}'."
        )

    return workbook


def workbook_response(
    workbook,
    filename,
):
    """
    Return an Excel workbook as a downloadable HTTP response.
    """

    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    return response