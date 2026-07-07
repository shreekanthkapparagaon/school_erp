from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.accounts.forms import PrincipalSchoolAssignmentForm
from apps.accounts.models import User
from apps.schools.forms import TeacherSubjectAssignmentForm
from apps.schools.models import School, SchoolClass, Subject, TeacherSubjectAssignment

from .models import Result
from .models import Exam


class ResultImportExportTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(
            name="Results School",
            code="RES",
            phone="1111111111",
            address="Results Road",
        )
        self.other_school = School.objects.create(
            name="Open School",
            code="OPEN",
            phone="2222222222",
            address="Open Road",
        )
        self.principal = self.make_user(
            "principal@results.test",
            User.Role.PRINCIPAL,
            school=self.school,
        )
        self.unassigned_principal = self.make_user(
            "open-principal@results.test",
            User.Role.PRINCIPAL,
        )
        self.head = self.make_user(
            "head@results.test",
            User.Role.HEAD,
            school=self.school,
        )
        self.other_head = self.make_user(
            "other-head@results.test",
            User.Role.HEAD,
            school=self.school,
        )
        self.teacher = self.make_user(
            "teacher@results.test",
            User.Role.TEACHER,
            school=self.school,
        )
        self.school_class = SchoolClass.objects.create(
            school=self.school,
            name="Grade 9",
            section="A",
        )
        self.other_class = SchoolClass.objects.create(
            school=self.school,
            name="Grade 10",
            section="B",
        )
        self.student = self.make_user(
            "student@results.test",
            User.Role.STUDENT,
            school=self.school,
            school_class=self.school_class,
        )
        self.subject = Subject.objects.create(
            school=self.school,
            name="Mathematics",
            code="MATH",
            created_by=self.head,
        )
        self.other_subject = Subject.objects.create(
            school=self.school,
            name="Science",
            code="SCI",
            created_by=self.other_head,
        )
        self.midterm_exam = Exam.objects.create(
            school=self.school,
            name="Midterm",
            academic_year="2026-27",
            exam_type="Summative",
            created_by=self.head,
        )
        self.final_exam = Exam.objects.create(
            school=self.school,
            name="Final",
            academic_year="2026-27",
            exam_type="Summative",
            created_by=self.head,
        )
        self.school_class.subjects.add(self.subject, self.other_subject)
        TeacherSubjectAssignment.objects.create(
            school=self.school,
            teacher=self.teacher,
            subject=self.subject,
            school_class=self.school_class,
            assigned_by=self.head,
        )

    def make_user(self, email, role, **extra_fields):
        return User.objects.create_user(
            email=email,
            password="StrongPass123!",
            role=role,
            **extra_fields,
        )

    def workbook_upload(self, marks):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            ["student_email", "subject_code", "marks_obtained", "max_marks"]
        )
        sheet.append([self.student.email, self.subject.code, marks, 100])
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            "results.xlsx",
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    def test_manager_assignment_lists_only_unassigned_records(self):
        form = PrincipalSchoolAssignmentForm()
        self.assertQuerySetEqual(
            form.fields["principal"].queryset,
            [self.unassigned_principal],
        )
        self.assertQuerySetEqual(
            form.fields["school"].queryset,
            [self.other_school],
        )

    def test_head_assignment_lists_only_subjects_created_by_head(self):
        form = TeacherSubjectAssignmentForm(manager=self.head)
        self.assertQuerySetEqual(
            form.fields["subject"].queryset,
            [self.subject],
        )
        principal_form = TeacherSubjectAssignmentForm(manager=self.principal)
        self.assertQuerySetEqual(
            principal_form.fields["subject"].queryset.order_by("code"),
            [self.subject, self.other_subject],
            ordered=False,
        )

    def test_teacher_class_list_and_roster_are_read_only(self):
        self.client.force_login(self.teacher)
        response = self.client.get(reverse("schools:class_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.school_class.name)
        self.assertNotContains(response, self.other_class.name)
        roster = self.client.get(
            reverse("schools:class_detail", args=[self.school_class.pk])
        )
        self.assertEqual(roster.status_code, 200)
        self.assertContains(roster, self.student.email)
        self.assertNotContains(roster, "Delete Class")

    def test_bulk_import_creates_then_updates_result(self):
        self.client.force_login(self.head)
        import_url = reverse(
            "results:import_results",
            args=[self.school_class.pk],
        )
        response = self.client.post(
            import_url,
            {
                "exam": self.midterm_exam.pk,
                "academic_year": "2026-27",
                "published": "on",
                "file": self.workbook_upload(72),
            },
        )
        self.assertEqual(response.status_code, 302)
        result = Result.objects.get()
        self.assertEqual(result.marks_obtained, 72)
        self.assertEqual(result.grade, "B")

        response = self.client.post(
            import_url,
            {
                "exam": self.midterm_exam.pk,
                "academic_year": "2026-27",
                "published": "on",
                "file": self.workbook_upload(91),
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Result.objects.count(), 1)
        result.refresh_from_db()
        self.assertEqual(result.marks_obtained, 91)
        self.assertEqual(result.grade, "A+")

    def test_principal_export_uses_separate_class_sheets(self):
        Result.objects.create(
            school=self.school,
            school_class=self.school_class,
            student=self.student,
            subject=self.subject,
            exam=self.final_exam,
            academic_year="2026-27",
            marks_obtained=88,
            max_marks=100,
            published=True,
            created_by=self.head,
        )
        self.client.force_login(self.principal)
        response = self.client.get(reverse("results:export_all"))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            set(workbook.sheetnames),
            {"Grade 9-A", "Grade 10-B"},
        )

        self.client.force_login(self.head)
        self.assertEqual(
            self.client.get(reverse("results:export_all")).status_code,
            302,
        )
