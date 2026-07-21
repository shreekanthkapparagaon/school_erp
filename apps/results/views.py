import re
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from apps.plans.utils import feature_action_required
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from tablib import Dataset

from apps.accounts.models import User
from apps.schools.models import SchoolClass, Subject, TeacherSubjectAssignment

from .forms import BulkResultImportForm, ResultForm
from .models import Exam, Result

EXPORT_HEADERS = [
    "Student",
    "Email",
    "Subject",
    "Subject Code",
    "Marks",
    "Maximum Marks",
    "Percentage",
    "Grade",
    "Exam",
    "Academic Year",
    "Published",
]


def can_manage_results(user, school_class):
    return bool(
        user.role in [User.Role.HEAD, User.Role.PRINCIPAL]
        and user.school_id == school_class.school_id
    )


def can_view_class_results(user, school_class):
    if can_manage_results(user, school_class):
        return True
    if user.role == User.Role.TEACHER:
        return school_class.teacher_assignments.filter(teacher=user).exists()
    return user.role == User.Role.STUDENT and user.school_class_id == school_class.pk


def can_view_result(user, result):
    if can_manage_results(user, result.school_class):
        return True
    if user.role == User.Role.TEACHER:
        return TeacherSubjectAssignment.objects.filter(
            teacher=user,
            school_class=result.school_class,
            subject=result.subject,
        ).exists()
    return (
        user.role == User.Role.STUDENT
        and result.student_id == user.pk
        and result.published
    )


def managed_classes(user):
    if user.role in [User.Role.HEAD, User.Role.PRINCIPAL] and user.school_id:
        return SchoolClass.objects.filter(school_id=user.school_id)
    if user.role == User.Role.TEACHER:
        return SchoolClass.objects.filter(
            teacher_assignments__teacher=user
        ).distinct()
    return SchoolClass.objects.none()


@login_required
@feature_action_required(
    "results",
    "view",
    "Your current subscription does not include Results.",
)
def result_home(request):
    if request.user.role == User.Role.STUDENT:
        if request.user.school_class_id:
            return redirect(
                "results:class_results",
                class_id=request.user.school_class_id,
            )
        messages.info(request, "No class has been assigned to your account.")
        return redirect("accounts:role_redirect")

    if request.user.role not in [
        User.Role.HEAD,
        User.Role.PRINCIPAL,
        User.Role.TEACHER,
    ]:
        messages.error(request, "You do not have permission to view results.")
        return redirect("accounts:role_redirect")

    exams = Exam.objects.filter(school_id=request.user.school_id).annotate(
        class_count=Count("results__school_class", distinct=True),
        result_count=Count("results", distinct=True),
    ).order_by("name")
    return render(
        request,
        "results/home.html",
        {
            "exams": exams,
            "can_manage": request.user.role in [User.Role.HEAD, User.Role.PRINCIPAL],
            "can_export_all": request.user.role == User.Role.PRINCIPAL,
        },
    )


@login_required
@feature_action_required(
    "results",
    "view",
    "Your current subscription does not include Results.",
)
def exam_results(request, exam_id):
    exam = get_object_or_404(
        Exam.objects.select_related("school"),
        pk=exam_id,
    )
    if request.user.role == User.Role.STUDENT:
        if request.user.school_class_id and request.user.school_id == exam.school_id:
            return redirect(
                "results:exam_class_results",
                exam_id=exam.pk,
                class_id=request.user.school_class_id,
            )
        messages.info(request, "No class has been assigned to your account.")
        return redirect("accounts:role_redirect")

    if request.user.role not in [
        User.Role.HEAD,
        User.Role.PRINCIPAL,
        User.Role.TEACHER,
    ]:
        messages.error(request, "You do not have permission to view results.")
        return redirect("accounts:role_redirect")

    if request.user.role == User.Role.TEACHER:
        classes = managed_classes(request.user).filter(school=exam.school)
    else:
        classes = SchoolClass.objects.filter(school=exam.school)

    classes = classes.annotate(
        result_count=Count("results", filter=Q(results__exam=exam))
    ).order_by("name")

    return render(
        request,
        "results/exam_results.html",
        {
            "exam": exam,
            "classes": classes,
            "can_manage": request.user.role in [User.Role.HEAD, User.Role.PRINCIPAL],
            "can_export_all": request.user.role == User.Role.PRINCIPAL,
        },
    )


@login_required
@feature_action_required(
    "results",
    "view",
    "You do not have permission to view results.",
)
def class_results(request, class_id, exam_id=None):
    school_class = get_object_or_404(
        SchoolClass.objects.select_related("school"),
        pk=class_id,
    )
    if not can_view_class_results(request.user, school_class):
        messages.error(request, "You do not have permission to view that class result.")
        return redirect("accounts:role_redirect")

    exams = Exam.objects.filter(school=school_class.school).order_by("name")
    exam = None
    if exam_id is not None:
        exam = get_object_or_404(Exam, pk=exam_id, school=school_class.school)

    results = Result.objects.filter(school_class=school_class)
    if exam is not None:
        results = results.filter(exam=exam)
    results = results.select_related("student", "subject", "school_class", "exam")
    if request.user.role == User.Role.STUDENT:
        results = results.filter(student=request.user, published=True)
    elif request.user.role == User.Role.TEACHER:
        subject_ids = TeacherSubjectAssignment.objects.filter(
            teacher=request.user,
            school_class=school_class,
        ).values_list("subject_id", flat=True)
        results = results.filter(subject_id__in=subject_ids)

    return render(
        request,
        "results/class_results.html",
        {
            "school_class": school_class,
            "exam": exam,
            "exams": exams,
            "results": results,
            "can_manage": can_manage_results(request.user, school_class),
            "can_export": request.user.role != User.Role.STUDENT,
        },
    )


@login_required
@feature_action_required(
    "results",
    "create",
    "You do not have permission to create results.",
)
def create_result(request, class_id, exam_id=None):
    school_class = get_object_or_404(SchoolClass, pk=class_id)
    if not can_manage_results(request.user, school_class):
        messages.error(request, "You do not have permission to create results.")
        return redirect("accounts:role_redirect")
    exam = None
    if exam_id is not None:
        exam = get_object_or_404(Exam, pk=exam_id, school=school_class.school)
    form = ResultForm(
        request.POST or None,
        school_class=school_class,
        manager=request.user,
        instance=Result(exam=exam) if exam else None,
    )
    if request.method == "POST" and form.is_valid():
        result = form.save()
        messages.success(request, "Result created.")
        return redirect("results:result_detail", result_id=result.pk)
    cancel_link = reverse("results:class_results", args=[school_class.pk])
    if exam is not None:
        cancel_link = reverse("results:exam_class_results", args=[exam.pk, school_class.pk])
    return render(
        request,
        "schools/form.html",
        {
            "form": form,
            "title": "Create Result",
            "submit_label": "Save result",
            "cancel_url": cancel_link,
        },
    )


@login_required
@feature_action_required(
    "results",
    "view",
    "You do not have permission to view result details.",
)
def result_detail(request, result_id):
    result = get_object_or_404(
        Result.objects.select_related(
            "school", "school_class", "student", "subject", "exam", "created_by"
        ),
        pk=result_id,
    )
    if not can_view_result(request.user, result):
        messages.error(request, "You do not have permission to view that result.")
        return redirect("accounts:role_redirect")
    can_edit = can_manage_results(request.user, result.school_class)
    return render(
        request,
        "shared/detail.html",
        {
            "title": str(result),
            "eyebrow": "Published Result" if result.published else "Draft Result",
            "fields": [
                ("Student", result.student),
                ("Class", result.school_class),
                ("Subject", result.subject),
                ("Exam", result.exam),
                ("Academic Year", result.academic_year),
                ("Marks", f"{result.marks_obtained:.2f} / {result.max_marks:.2f}"),
                ("Percentage", f"{result.percentage:.2f}%"),
                ("Grade", result.grade),
            ],
            "back_url": reverse(
                "results:class_results",
                args=[result.school_class_id],
            ),
            "edit_url": reverse("results:update_result", args=[result.pk]) if can_edit else "",
            "delete_url": reverse("results:delete_result", args=[result.pk]) if can_edit else "",
        },
    )


@login_required
@feature_action_required(
    "results",
    "update",
    "You do not have permission to update results.",
)
def update_result(request, result_id):
    result = get_object_or_404(
        Result.objects.select_related("school_class", "school", "exam"),
        pk=result_id,
    )
    if not can_manage_results(request.user, result.school_class):
        messages.error(request, "You do not have permission to edit that result.")
        return redirect("accounts:role_redirect")
    form = ResultForm(
        request.POST or None,
        instance=result,
        school_class=result.school_class,
        manager=request.user,
    )
    if request.method == "POST" and form.is_valid():
        result = form.save(commit=False)
        if "grade" not in form.changed_data:
            result.grade = ""
        result.save()
        messages.success(request, "Result updated.")
        return redirect("results:result_detail", result_id=result.pk)
    return render(
        request,
        "schools/form.html",
        {
            "form": form,
            "title": "Edit Result",
            "submit_label": "Save changes",
            "cancel_url": reverse("results:result_detail", args=[result.pk]),
        },
    )


@login_required
@feature_action_required(
    "results",
    "delete",
    "You do not have permission to delete results.",
)
def delete_result(request, result_id):
    result = get_object_or_404(
        Result.objects.select_related("school_class", "student", "subject"),
        pk=result_id,
    )
    if not can_manage_results(request.user, result.school_class):
        messages.error(request, "You do not have permission to delete that result.")
        return redirect("accounts:role_redirect")
    class_id = result.school_class_id
    if request.method == "POST":
        label = str(result)
        result.delete()
        messages.success(request, f"Deleted {label}.")
        return redirect("results:class_results", class_id=class_id)
    return render(
        request,
        "shared/confirm_delete.html",
        {
            "title": "Delete Result",
            "object_label": str(result),
            "warning": "This permanently removes the student's result record.",
            "cancel_url": reverse("results:result_detail", args=[result.pk]),
        },
    )


@login_required
@feature_action_required(
    "results",
    "import",
    "You do not have permission to import results.",
)
def import_results(request, class_id):
    school_class = get_object_or_404(
        SchoolClass.objects.select_related("school").prefetch_related("subjects"),
        pk=class_id,
    )
    if not can_manage_results(request.user, school_class):
        messages.error(request, "You do not have permission to import results.")
        return redirect("accounts:role_redirect")

    form = BulkResultImportForm(
        request.POST or None,
        request.FILES or None,
        school_class=school_class,
    )
    if request.method == "POST" and form.is_valid():
        try:
            dataset = Dataset().load(form.cleaned_data["file"].read(), format="xlsx")
            required = {
                "student_email",
                "subject_code",
                "marks_obtained",
                "max_marks",
            }
            missing = required.difference(dataset.headers or [])
            if missing:
                raise ValueError(f"Missing columns: {', '.join(sorted(missing))}")

            imported = 0
            with transaction.atomic():
                for row_number, row in enumerate(dataset.dict, start=2):
                    if not any(row.get(column) not in [None, ""] for column in required):
                        continue
                    student = User.objects.get(
                        email__iexact=str(row["student_email"]).strip(),
                        role=User.Role.STUDENT,
                        school_class=school_class,
                        school=school_class.school,
                    )
                    subject = Subject.objects.get(
                        school=school_class.school,
                        classes=school_class,
                        code__iexact=str(row["subject_code"]).strip(),
                    )
                    try:
                        marks = Decimal(str(row["marks_obtained"]))
                        maximum = Decimal(str(row["max_marks"]))
                    except (InvalidOperation, TypeError):
                        raise ValueError(
                            f"Row {row_number}: marks and maximum marks must be numbers."
                        )

                    exam = form.cleaned_data["exam"]
                    result = Result.objects.filter(
                        school_class=school_class,
                        student=student,
                        subject=subject,
                        exam=exam,
                    ).first() or Result(
                        school=school_class.school,
                        school_class=school_class,
                        student=student,
                        subject=subject,
                        exam=exam,
                        created_by=request.user,
                    )
                    result.marks_obtained = marks
                    result.max_marks = maximum
                    result.grade = ""
                    result.published = form.cleaned_data["published"]
                    result.full_clean()
                    result.save()
                    imported += 1
        except (ValueError, ValidationError, User.DoesNotExist, Subject.DoesNotExist) as exc:
            form.add_error("file", str(exc))
        else:
            messages.success(request, f"Imported or updated {imported} results.")
            return redirect("results:class_results", class_id=school_class.pk)

    return render(
        request,
        "results/import.html",
        {"form": form, "school_class": school_class},
    )


def style_sheet(sheet):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[letter].width = min(max(width + 2, 12), 34)


def safe_sheet_title(value, used_titles):
    title = re.sub(r"[\[\]:*?/\\]", "-", value).strip()[:31] or "Class"
    original = title
    number = 2
    while title in used_titles:
        suffix = f"-{number}"
        title = f"{original[:31-len(suffix)]}{suffix}"
        number += 1
    used_titles.add(title)
    return title


def append_results_sheet(workbook, school_class, results, used_titles):
    title = safe_sheet_title(
        f"{school_class.name}-{school_class.section}".strip("-"),
        used_titles,
    )
    sheet = workbook.create_sheet(title)
    sheet.append(EXPORT_HEADERS)
    for result in results:
        sheet.append(
            [
                str(result.student),
                result.student.email,
                result.subject.name,
                result.subject.code,
                float(result.marks_obtained),
                float(result.max_marks),
                float(result.percentage),
                result.grade,
                result.exam_name,
                result.academic_year,
                "Yes" if result.published else "No",
            ]
        )
    style_sheet(sheet)


def workbook_response(workbook, filename):
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
@feature_action_required(
    "results",
    "import",
    "You do not have permission to download the import template.",
)
def download_import_template(request, class_id):
    school_class = get_object_or_404(
        SchoolClass.objects.select_related("school").prefetch_related("subjects"),
        pk=class_id,
    )
    if not can_manage_results(request.user, school_class):
        messages.error(request, "You do not have permission to download that template.")
        return redirect("accounts:role_redirect")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results Import"
    sheet.append(["student_email", "subject_code", "marks_obtained", "max_marks"])
    students = school_class.students.order_by("email")
    subjects = school_class.subjects.order_by("code")
    for student in students:
        for subject in subjects:
            sheet.append([student.email, subject.code, "", 100])
    style_sheet(sheet)
    return workbook_response(
        workbook,
        f"{school_class.name}-{school_class.section}-result-template.xlsx",
    )


@login_required
@feature_action_required(
    "results",
    "export",
    "You do not have permission to export results.",
)
def export_class_results(request, class_id):
    school_class = get_object_or_404(
        SchoolClass.objects.select_related("school"),
        pk=class_id,
    )
    if not can_view_class_results(request.user, school_class) or request.user.role == User.Role.STUDENT:
        messages.error(request, "You do not have permission to export that class result.")
        return redirect("accounts:role_redirect")

    results = Result.objects.filter(school_class=school_class).select_related(
        "student", "subject"
    )
    if request.user.role == User.Role.TEACHER:
        subject_ids = TeacherSubjectAssignment.objects.filter(
            teacher=request.user,
            school_class=school_class,
        ).values_list("subject_id", flat=True)
        results = results.filter(subject_id__in=subject_ids)

    workbook = Workbook()
    workbook.remove(workbook.active)
    append_results_sheet(workbook, school_class, results, set())
    return workbook_response(
        workbook,
        f"{school_class.name}-{school_class.section}-results.xlsx",
    )


@login_required
@feature_action_required(
    "results",
    "export",
    "You do not have permission to export results.",
)
def export_all_results(request):
    if request.user.role != User.Role.PRINCIPAL or not request.user.school_id:
        messages.error(request, "Only the school principal can export all class results.")
        return redirect("accounts:role_redirect")

    classes = SchoolClass.objects.filter(school_id=request.user.school_id).order_by(
        "name", "section"
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles = set()
    for school_class in classes:
        results = Result.objects.filter(school_class=school_class).select_related(
            "student", "subject"
        )
        append_results_sheet(workbook, school_class, results, used_titles)

    if not workbook.worksheets:
        sheet = workbook.create_sheet("Results")
        sheet.append(EXPORT_HEADERS)
        style_sheet(sheet)
    return workbook_response(
        workbook,
        f"{request.user.school.code}-all-class-results.xlsx",
    )
